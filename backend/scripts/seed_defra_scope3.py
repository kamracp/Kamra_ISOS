"""Seed DEFRA 2026 Scope-3 factors into emission_factors (idempotent; --dry to preview).

  WTT- fuels       (tonnes rows)            -> meter_type wtt_<fuel>,               unit tonne
  Freighting goods (tonne.km, diesel/none)  -> meter_type freight_<mode>_<class>,   unit tonne.km
  Waste disposal   (tonnes x treatment)     -> meter_type waste_<material>_<route>, unit tonne
  WTT- bioenergy   (biomass/biogas, tonnes)  -> meter_type wtt_<fuel>,  unit tonne
  Business travel- land/air/sea + WTT counterparts (km / passenger.km, no miles, no market-segment cars)
                                            -> meter_type pass_<mode>_... / wtt_pass_<mode>_..., unit km or passenger.km; cat 6/7

Same conventions as seed_defra_materials.py: region UK, source_year 2026, DEFRA ID in
document_reference, skip when meter_type+source_year already present.
Usage: PGPASSWORD=postgres python scripts/seed_defra_scope3.py [--dry]
"""
import glob
import os
import re
import sys

import psycopg2
from openpyxl import load_workbook

SRC = "DEFRA 2026 GHG Conversion Factors (flat file v1.2, revised 31 Jul 2026)"
FREIGHT_MODES = {"HGV (non-refrigerated, all diesel)", "Rail", "Sea tanker", "Cargo ship", "Freight flights"}
# Passenger sections -> meter_type prefix. Vehicle-only and WTT are seeded as SEPARATE factors;
# the user enters one line per factor (the engine never sums them silently).
PASS_L1 = {"Business travel- land": "pass", "Business travel- air": "pass", "Business travel- sea": "pass",
           "WTT- pass vehs & travel- land": "wtt_pass", "WTT- business travel- air": "wtt_pass", "WTT- business travel- sea": "wtt_pass"}
PASS_MODES = {"cars (by size)": "car", "motorbike": "motorbike", "taxis": "taxi", "bus": "bus", "rail": "rail", "flights": "flight", "ferry": "ferry"}


def slug(*parts: str) -> str:
    s = "_".join(str(p) for p in parts if p)
    s = re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")
    return s


def rows_from_flat() -> list[dict]:
    ws = load_workbook(glob.glob("data/defra/*.xlsx")[0], read_only=True)["Factors by Category"]
    it = ws.iter_rows(min_row=6, values_only=True)
    hdr = next(it); H = {h: i for i, h in enumerate(hdr)}
    out = []
    for r in it:
        if r[H["Scope"]] != "Scope 3" or r[H["GHG/Unit"]] != "kg CO2e" or r[H["GHG Conversion Factor 2026"]] is None:
            continue
        l1, l2, l3, ct, uom = r[H["Level 1"]], r[H["Level 2"]], r[H["Level 3"]], r[H["Column Text"]], r[H["UOM"]]
        val, defra_id = float(r[H["GHG Conversion Factor 2026"]]), r[H["ID"]]
        l3 = (l3 or "").strip()
        if l1 == "WTT- fuels" and uom == "tonnes":
            mt, unit, note = slug("wtt", l3), "tonne", f"WTT (well-to-tank) upstream factor for {l3}; Scope 3 cat. 3"
        elif l1 == "Freighting goods" and uom == "tonne.km" and l2 in FREIGHT_MODES and ct in (None, "Diesel", "Average laden", "100% Laden"):
            laden = "" if ct in (None, "Diesel") else ("avg_laden" if ct == "Average laden" else "full_laden")
            l2s = "hgv" if l2.startswith("HGV (non") else l2
            mt, unit, note = slug("freight", l2s, l3, laden), "tonne.km", f"Freight {l2} / {l3}" + (f" / {ct}" if laden else "") + "; Scope 3 cat. 4/9 (vehicle-only, excl. WTT)"
        elif l1 == "WTT- bioenergy" and l2 in ("WTT- biomass", "WTT- biogas") and uom == "tonnes":
            mt, unit, note = slug("wtt", l3), "tonne", f"WTT (well-to-tank) upstream factor for {l3} (DEFRA {l2}); Scope 3 cat. 3"
        elif l1 in PASS_L1 and uom in ("km", "passenger.km") and str(l2).replace("WTT- ", "").lower() in PASS_MODES:
            mode = PASS_MODES[str(l2).replace("WTT- ", "").lower()]
            l4 = (r[H["Level 4"]] or "").strip() if "Level 4" in H else ""
            # taxis publish both km and passenger.km rows with identical names -> suffix the passenger.km one
            mt = slug(PASS_L1[l1], mode, l3, l4, ct) + ("_pkm" if mode == "taxi" and uom == "passenger.km" else "")
            unit = uom
            note = (f"Passenger {l2} / {l3}" + (f" / {l4}" if l4 else "") + (f" / {ct}" if ct else "") + "; Scope 3 cat. 6/7"
                    + (" (WTT upstream; enter as a separate line to the vehicle factor)" if PASS_L1[l1] == "wtt_pass" else " (vehicle-only, excl. WTT)"))
        elif l1 == "Waste disposal" and uom == "tonnes":
            mt, unit, note = slug("waste", l3, ct), "tonne", f"Waste treatment: {l3} via {ct}; Scope 3 cat. 5"
        else:
            continue
        out.append({"mt": mt, "unit": unit, "val": val, "ref": f"DEFRA 2026 ID {defra_id}: {l1} / {l2} / {l3}" + (f" / {ct}" if ct else ""), "note": note})
    # first occurrence wins when several flat rows collapse to one meter_type
    seen, uniq = set(), []
    for o in out:
        if o["mt"] not in seen:
            seen.add(o["mt"]); uniq.append(o)
    return uniq


def main() -> None:
    dry = "--dry" in sys.argv
    rows = rows_from_flat()
    print(f"candidates: {len(rows)}  (wtt={sum(r['mt'].startswith('wtt_') and not r['mt'].startswith('wtt_pass_') for r in rows)}, "
          f"pass={sum(r['mt'].startswith(('pass_', 'wtt_pass_')) for r in rows)}, "
          f"freight={sum(r['mt'].startswith('freight_') for r in rows)}, waste={sum(r['mt'].startswith('waste_') for r in rows)})")
    conn = psycopg2.connect(host="127.0.0.1", user="postgres", password=os.environ.get("PGPASSWORD", "postgres"), dbname="benas")
    cur = conn.cursor(); ins = skip = 0
    for r in rows:
        cur.execute("SELECT 1 FROM emission_factors WHERE meter_type=%s AND source_year=2026", (r["mt"],))
        if cur.fetchone():
            skip += 1; continue
        if not dry:
            cur.execute("INSERT INTO emission_factors (meter_type, unit, factor_kgco2e_per_unit, region, source, source_year, "
                        "document_reference, valid_from, is_active, notes) VALUES (%s,%s,%s,'UK',%s,2026,%s,'2026-01-01',true,%s)",
                        (r["mt"], r["unit"], r["val"], SRC, r["ref"], r["note"]))
        ins += 1
    conn.commit(); print(f"inserted={ins} skipped={skip} dry={dry}")
    if dry:
        for r in rows[:6]: print("  e.g.", r["mt"], r["unit"], r["val"])


if __name__ == "__main__":
    main()
