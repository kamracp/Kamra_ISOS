"""List DEFRA 2026 Material use rows for Primary material production only."""
from openpyxl import load_workbook
ws = load_workbook("data/defra/defra-2026-flat.xlsx", read_only=True, data_only=True)["Factors by Category"]
rows = list(ws.iter_rows(min_row=6, values_only=True))
h = [str(c).strip() if c else "" for c in rows[0]]
c = {n: h.index(n) for n in h if n}
for r in rows[1:]:
    if r and r[c["Level 1"]] == "Material use" and str(r[c["Column Text"]]).startswith("Primary"):
        print(r[c["ID"]], "|", r[c["Level 2"]], "|", r[c["Level 3"]], "|", r[c["Column Text"]], "|", r[c["UOM"]], "|", r[c["GHG Conversion Factor 2026"]])
