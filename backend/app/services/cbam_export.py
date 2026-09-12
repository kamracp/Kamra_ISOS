"""CBAM operator -> importer communication workbook (XLSX).

Structured on the data elements of Implementing Regulation (EU) 2023/1773 Annex IV
(the layout EU importers know from the Commission's transitional 'Summary_Communication'
sheet), with SEE computed per Implementing Regulation (EU) 2025/2547. This is Kamra's
own workbook, not the Commission file: the Commission has published no definitive-period
operator template as of Sep 2026. Pure function -- no DB access.
"""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_MASS_TO_TONNE = {"tonne": 1.0, "t": 1.0, "mt": 1.0, "kg": 0.001}
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E78")
SECTION = Font(bold=True, size=12)


def _header(ws, row: int, cols: list[str]) -> None:
    for c, name in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=c, value=name)
        cell.font, cell.fill = HEAD, FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _autowidth(ws, minimum: int = 12, maximum: int = 60) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(minimum, min(maximum, width + 2))


def _electricity_per_tonne(product: dict | None) -> tuple[float | None, float | None, str | None]:
    """MWh per tonne of good, grid factor (tCO2e/MWh) and its source, from the LCA items that
    count as indirect (grid items not overridden away). None when the product is unusable."""
    if not product:
        return None, None, None
    fu = _MASS_TO_TONNE.get((product.get("functional_unit") or "").strip().lower())
    if fu is None:
        return None, None, None
    tonnes_per_fu = (product.get("functional_unit_qty") or 0.0) * fu
    if tonnes_per_fu <= 0:
        return None, None, None
    kwh, factor, source = 0.0, None, None
    for i in product.get("items", []):
        bucket = i.get("cbam_bucket") or ("indirect" if i.get("factor_source") == "electricity" else "")
        if i.get("factor_source") == "electricity" and bucket == "indirect" and i.get("status") == "calculated":
            kwh += i.get("quantity_per_fu") or 0.0
            factor, source = i.get("factor_value"), i.get("factor_citation")
    if kwh == 0.0:
        return None, None, None
    return round(kwh / 1000.0 / tonnes_per_fu, 6), (round(factor, 6) if factor is not None else None), source


def build_operator_template_xlsx(org, goods: list[dict], products: dict[int, dict], year: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary_Communication"
    ws["A1"], ws["A1"].font = "CBAM - Communication of embedded emissions from operator to importer", SECTION
    ws["A2"] = f"Reporting period: calendar year {year}   |   Generated: {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC   |   Kamra ClimateOS"
    ws["A4"], ws["A4"].font = "A. Installation / operator", SECTION
    rows_a = [("Operator (legal name)", getattr(org, "legal_name", None) or getattr(org, "organization_name", None)),
              ("Operator (trading name)", getattr(org, "organization_name", None)),
              ("Country", getattr(org, "country", None)), ("State / city", f"{getattr(org, 'state', '') or ''} / {getattr(org, 'city', '') or ''}"),
              ("Contact e-mail", getattr(org, "email", None)), ("Website", getattr(org, "website", None))]
    for r, (k, v) in enumerate(rows_a, start=5):
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)

    start = 5 + len(rows_a) + 1
    ws.cell(row=start, column=1, value="B. Goods and specific embedded emissions (SEE)").font = SECTION
    cols = ["Good", "CN code", "CBAM category", "Production route", "Reporting period",
            "Activity level (t produced)", "SEE direct (tCO2e/t)", "SEE indirect (tCO2e/t)",
            "Indirect counted in total?", "SEE total (tCO2e/t)", "Total embedded (tCO2e)",
            "Electricity (MWh/t)", "Grid EF (tCO2e/MWh)", "Grid EF source",
            "Carbon price paid (per tCO2e)", "Calculation status", "Status / note"]
    _header(ws, start + 1, cols)
    for r, g in enumerate(goods, start=start + 2):
        mwh, ef, src = _electricity_per_tonne(products.get(g.get("lca_product_id") or -1))
        vals = [g["name"], g["cn_code"], g["goods_category"], g.get("production_route"), f"{year}-01-01 to {year}-12-31",
                g["production_qty_tonne"], g["see_direct_tco2e_per_t"], g["see_indirect_tco2e_per_t"],
                "yes" if g["indirect_required"] else "no (Annex II good)", g["see_total_tco2e_per_t"], g["total_embedded_tco2e"],
                mwh, ef, src, g.get("carbon_price_paid_per_tco2e"), g["status"], g.get("status_reason")]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)
    ws.freeze_panes = ws.cell(row=start + 2, column=2)
    _autowidth(ws)
    ws.column_dimensions["A"].width = 34   # title rows are long; do not let them stretch column A

    wp = wb.create_sheet("Precursors")
    _header(wp, 1, ["Good", "CN code", "Inventory item", "Precursor good", "Precursor CN", "Precursor per t of good (t/t)",
                    "Precursor SEE direct (tCO2e/t)", "Precursor SEE indirect (tCO2e/t)", "Direct contributed (kg CO2e per FU)"])
    r = 2
    for g in goods:
        for p in g.get("precursors", []):
            for c, v in enumerate([g["name"], g["cn_code"], p["item_name"], p["precursor_name"], p["cn_code"], p["tonne_per_fu"],
                                   p["see_direct_tco2e_per_t"], p["see_indirect_tco2e_per_t"], p["direct_kg_per_fu"]], start=1):
                wp.cell(row=r, column=c, value=v)
            r += 1
    if r == 2:
        wp.cell(row=2, column=1, value="No precursor goods linked for this period.")
    _autowidth(wp)

    wm = wb.create_sheet("Method_Sources")
    wm["A1"], wm["A1"].font = "Methodology and data sources", SECTION
    notes = [
        "Method: SEE = attributed emissions of the production process / activity level (Regulation (EU) 2023/956 Annex IV; "
        "Implementing Regulation (EU) 2025/2547 on calculation methods). Direct = fuel combustion + process emissions + precursor "
        "embedded emissions; indirect = purchased electricity x grid emission factor.",
        "Emission factors are resolved at read time from cited sources (IPCC 2006 Vol.2/Vol.3, CEA India grid database, DEFRA "
        "conversion factors); no hand-typed or estimated factor values are used.",
        "Layout follows the data elements of Implementing Regulation (EU) 2023/1773 Annex IV. This workbook is generated by Kamra "
        "ClimateOS; it is not the European Commission's file.",
    ]
    for i, n in enumerate(notes, start=2):
        wm.cell(row=i, column=1, value=n).alignment = Alignment(wrap_text=True)
    _header(wm, 6, ["Good", "Factor sources used", "Items outside CBAM boundary / not counted"])
    for r, g in enumerate(goods, start=7):
        wm.cell(row=r, column=1, value=g["name"])
        wm.cell(row=r, column=2, value=" | ".join(g.get("factor_sources", []))).alignment = Alignment(wrap_text=True)
        wm.cell(row=r, column=3, value="\n".join(g.get("excluded_items", [])) or "-").alignment = Alignment(wrap_text=True)
    wm.column_dimensions["A"].width, wm.column_dimensions["B"].width, wm.column_dimensions["C"].width = 30, 90, 70

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
